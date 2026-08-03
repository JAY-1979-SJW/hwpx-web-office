"""HWPX 서식 카탈로그(data/drafts/form_library/catalog.sqlite) 목차 생성.

읽기 전용 - 원본 카탈로그(38,165건, §4.4 읽기 전용 구역)는 절대 안 건드리고
목차만 새로 뽑는다. 산출: data/reports/hwpx_form_catalog_toc/
  - toc_summary.md          : 발행기관·문서유형·서식종류별 건수 요약(사람이 읽는 것)
  - hwpx_form_catalog_toc.xlsx : 요약 시트 + 전체 38,165건 상세 시트
  - catalog_toc_viewer.html : 발행기관/문서유형/서식종류 사이드 메뉴로
    거르는 자체완결형(서버 불필요, JSON 인라인) 뷰어

사용: python scripts/ops/build_hwpx_form_catalog_toc.py
"""
from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[2]
DB_PATH = ROOT / "data" / "drafts" / "form_library" / "catalog.sqlite"
OUT_DIR = ROOT / "data" / "reports" / "hwpx_form_catalog_toc"


def _conn() -> sqlite3.Connection:
    conn = sqlite3.connect(f"file:{DB_PATH}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    return conn


def _fetch_rows(c: sqlite3.Connection) -> list[sqlite3.Row]:
    return c.execute(
        """SELECT f.form_id, f.name, f.institution, f.legal_basis, f.source_path,
                  d.clean_name, d.doc_type, d.form_kind, d.fillable, d.input_count
           FROM forms f LEFT JOIN derivations_rebuild d ON f.form_id = d.form_id
           ORDER BY f.institution, d.doc_type, f.form_id"""
    ).fetchall()


def _group_counts(rows: list[sqlite3.Row], key: str) -> list[tuple[str, int]]:
    counts: dict[str, int] = {}
    for r in rows:
        k = (r[key] or "(미분류)").strip() or "(미분류)"
        counts[k] = counts.get(k, 0) + 1
    return sorted(counts.items(), key=lambda x: -x[1])


def build() -> dict:
    if not DB_PATH.exists():
        raise SystemExit(f"카탈로그 DB 없음: {DB_PATH}")
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    with _conn() as c:
        rows = _fetch_rows(c)

    by_inst = _group_counts(rows, "institution")
    by_doc = _group_counts(rows, "doc_type")
    by_kind = _group_counts(rows, "form_kind")
    fillable_n = sum(1 for r in rows if r["fillable"])

    # ── xlsx: 요약 시트 + 상세 시트 ──────────────────────────────────────
    wb = openpyxl.Workbook()
    ws_sum = wb.active
    ws_sum.title = "요약"
    ws_sum.append(["구분", "값"])
    ws_sum["A1"].font = ws_sum["B1"].font = Font(bold=True)
    ws_sum.append(["총 서식 건수", len(rows)])
    ws_sum.append(["입력칸 있는 서식(fillable)", fillable_n])
    ws_sum.append([])
    ws_sum.append(["발행기관", "건수"])
    for k, n in by_inst[:40]:
        ws_sum.append([k, n])
    r0 = ws_sum.max_row + 2
    ws_sum.cell(row=r0, column=1, value="문서유형(doc_type)")
    ws_sum.cell(row=r0, column=2, value="건수")
    for i, (k, n) in enumerate(by_doc, start=1):
        ws_sum.cell(row=r0 + i, column=1, value=k)
        ws_sum.cell(row=r0 + i, column=2, value=n)
    r1 = ws_sum.max_row + 2
    ws_sum.cell(row=r1, column=1, value="서식종류(form_kind)")
    ws_sum.cell(row=r1, column=2, value="건수")
    for i, (k, n) in enumerate(by_kind, start=1):
        ws_sum.cell(row=r1 + i, column=1, value=k)
        ws_sum.cell(row=r1 + i, column=2, value=n)
    ws_sum.column_dimensions["A"].width = 40
    ws_sum.column_dimensions["B"].width = 12

    ws_detail = wb.create_sheet("전체 목차")
    header = ["form_id", "서식명", "발행기관", "문서유형", "서식종류", "입력가능", "입력칸수", "법적근거"]
    ws_detail.append(header)
    for cell in ws_detail[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws_detail.append([
            r["form_id"], r["clean_name"] or r["name"], r["institution"] or "",
            r["doc_type"] or "", r["form_kind"] or "",
            "Y" if r["fillable"] else "N", r["input_count"] or 0,
            (r["legal_basis"] or "")[:200],
        ])
    widths = {"A": 10, "B": 52, "C": 16, "D": 12, "E": 12, "F": 10, "G": 10, "H": 40}
    for col, w in widths.items():
        ws_detail.column_dimensions[col].width = w
    ws_detail.freeze_panes = "A2"

    xlsx_path = OUT_DIR / "hwpx_form_catalog_toc.xlsx"
    wb.save(xlsx_path)

    # ── 요약 md(사람이 채팅창 밖에서 훑어볼 것) ─────────────────────────
    md_lines = [
        "# HWPX 서식 카탈로그 목차 요약",
        "",
        f"- 총 서식: {len(rows)}건 (입력칸 있는 서식 {fillable_n}건)",
        "",
        "## 발행기관 상위 15",
    ]
    for k, n in by_inst[:15]:
        md_lines.append(f"- {k}: {n}건")
    md_lines.append("")
    md_lines.append("## 문서유형(doc_type)")
    for k, n in by_doc:
        md_lines.append(f"- {k}: {n}건")
    md_lines.append("")
    md_lines.append("## 서식종류(form_kind)")
    for k, n in by_kind:
        md_lines.append(f"- {k}: {n}건")
    (OUT_DIR / "toc_summary.md").write_text("\n".join(md_lines) + "\n", encoding="utf-8")

    # ── 사이드 메뉴 뷰어(html, JSON 인라인 - 서버 없이 바로 열림) ─────────
    items = [
        {
            "id": r["form_id"], "name": r["clean_name"] or r["name"],
            "inst": (r["institution"] or "(미분류)").strip() or "(미분류)",
            "doc": (r["doc_type"] or "(미분류)").strip() or "(미분류)",
            "kind": (r["form_kind"] or "(미분류)").strip() or "(미분류)",
            "fillable": bool(r["fillable"]), "inputCount": r["input_count"] or 0,
            "sourcePath": r["source_path"] or "",
        }
        for r in rows
    ]
    viewer_html = _render_viewer_html(items, by_inst, by_doc, by_kind, len(rows), fillable_n)
    viewer_path = OUT_DIR / "catalog_toc_viewer.html"
    viewer_path.write_text(viewer_html, encoding="utf-8")

    return {"total": len(rows), "fillable": fillable_n, "xlsx": str(xlsx_path),
            "viewer": str(viewer_path), "by_inst_top5": by_inst[:5], "by_doc": by_doc}


def _render_viewer_html(items, by_inst, by_doc, by_kind, total, fillable_n) -> str:
    data_json = json.dumps(items, ensure_ascii=False)
    groups_json = json.dumps({"발행기관": by_inst, "문서유형": by_doc, "서식종류": by_kind}, ensure_ascii=False)
    return """<!DOCTYPE html>
<html lang="ko"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>HWPX 서식 카탈로그 목차</title>
<style>
  :root {
    --bg: #f6f7f9; --panel: #ffffff; --border: #e4e7ec; --border-strong: #d7dbe3;
    --text: #1a1d23; --text-dim: #6b7280; --text-faint: #9aa1ad;
    --side-bg: #14181f; --side-text: #cbd2dc; --side-text-dim: #7c8494; --side-hover: #1f2530;
    --accent: #3457d5; --accent-soft: #eaeefc; --accent-text: #2643b0;
    --ok: #1b7f4d; --ok-soft: #e7f6ee; --no: #9aa1ad;
    --radius: 8px;
    --font-ui: "Pretendard", "Malgun Gothic", -apple-system, "Segoe UI", Roboto, sans-serif;
    --font-mono: "SFMono-Regular", Consolas, "Liberation Mono", monospace;
  }
  @media (prefers-color-scheme: dark) {
    :root {
      --bg: #14161a; --panel: #1b1e24; --border: #2a2e37; --border-strong: #383e4a;
      --text: #e7e9ee; --text-dim: #9aa1ad; --text-faint: #6b7280;
      --side-bg: #0c0e12; --side-text: #c3cad6; --side-text-dim: #6b7280; --side-hover: #191c22;
      --accent: #6d8bff; --accent-soft: #202a4a; --accent-text: #b7c4ff;
      --ok: #4fd08a; --ok-soft: #16321f; --no: #6b7280;
    }
  }
  :root[data-theme="dark"] {
    --bg: #14161a; --panel: #1b1e24; --border: #2a2e37; --border-strong: #383e4a;
    --text: #e7e9ee; --text-dim: #9aa1ad; --text-faint: #6b7280;
    --side-bg: #0c0e12; --side-text: #c3cad6; --side-text-dim: #6b7280; --side-hover: #191c22;
    --accent: #6d8bff; --accent-soft: #202a4a; --accent-text: #b7c4ff;
    --ok: #4fd08a; --ok-soft: #16321f; --no: #6b7280;
  }
  :root[data-theme="light"] {
    --bg: #f6f7f9; --panel: #ffffff; --border: #e4e7ec; --border-strong: #d7dbe3;
    --text: #1a1d23; --text-dim: #6b7280; --text-faint: #9aa1ad;
    --side-bg: #14181f; --side-text: #cbd2dc; --side-text-dim: #7c8494; --side-hover: #1f2530;
    --accent: #3457d5; --accent-soft: #eaeefc; --accent-text: #2643b0;
    --ok: #1b7f4d; --ok-soft: #e7f6ee; --no: #9aa1ad;
  }
  * { box-sizing: border-box; }
  body { margin: 0; font-family: var(--font-ui); display: flex; height: 100vh; background: var(--bg); color: var(--text); }
  ::-webkit-scrollbar { width: 10px; height: 10px; }
  ::-webkit-scrollbar-thumb { background: var(--border-strong); border-radius: 6px; }

  #sidebar { width: 280px; flex: none; background: var(--side-bg); color: var(--side-text);
             display: flex; flex-direction: column; padding: 18px 0 0; }
  #sidebar .brand { padding: 0 18px 16px; font-weight: 700; font-size: 0.95em; color: #fff;
                     display: flex; align-items: center; gap: 8px; }
  #sidebar .brand .dot { width: 8px; height: 8px; border-radius: 50%; background: var(--accent); flex: none; }
  #axis-tabs { display: flex; gap: 2px; padding: 0 12px 10px; }
  #axis-tabs button { flex: 1; font: inherit; font-size: 0.78em; font-weight: 600; padding: 7px 4px;
                       border: none; border-radius: 6px; background: transparent; color: var(--side-text-dim); cursor: pointer;
                       position: relative; }
  #axis-tabs button:hover { background: var(--side-hover); color: var(--side-text); }
  #axis-tabs button.on { background: var(--accent); color: #fff; }
  #axis-tabs button .n { position: absolute; top: -4px; right: -2px; background: var(--ok); color: #fff;
                          font-size: 0.7em; min-width: 14px; height: 14px; border-radius: 7px; line-height: 14px;
                          padding: 0 3px; }
  #axis-list { flex: 1; overflow-y: auto; padding-bottom: 18px; border-top: 1px solid #232833; }
  #sidebar .item { padding: 7px 18px; cursor: pointer; font-size: 0.87em; line-height: 1.5;
                    display: flex; align-items: center; gap: 9px; }
  #sidebar .item:hover { background: var(--side-hover); }
  #sidebar .item input[type="checkbox"] { accent-color: var(--accent); flex: none; width: 14px; height: 14px; }
  #sidebar .item .lbl { flex: 1; overflow: hidden; text-overflow: ellipsis; white-space: nowrap; }
  #sidebar .item.checked .lbl { color: #fff; font-weight: 600; }
  #sidebar .item.zero { opacity: 0.4; }
  #sidebar .cnt { color: var(--side-text-dim); font-variant-numeric: tabular-nums; flex: none; font-size: 0.92em; }
  #sidebar .item.checked .cnt { color: var(--accent-text); }
  #axis-clear { padding: 6px 18px 10px; }
  #axis-clear button { font: inherit; font-size: 0.78em; color: var(--accent-text); background: none;
                        border: none; cursor: pointer; padding: 0; }
  #axis-clear button:hover { text-decoration: underline; }

  #main { flex: 1; overflow-y: auto; padding: 28px 32px 40px; min-width: 0; }
  h1 { margin: 0 0 4px; font-size: 1.4em; font-weight: 700; letter-spacing: -.01em; }
  .sub { margin: 0 0 20px; color: var(--text-dim); font-size: 0.88em; }

  #toolbar { display: flex; gap: 10px; align-items: center; margin-bottom: 14px; }
  input#q, input#base, select#sort { font: inherit; padding: 8px 12px; border: 1px solid var(--border-strong);
                          border-radius: var(--radius); background: var(--panel); color: var(--text); }
  input#q { flex: 1; max-width: 380px; }
  input#q:focus, input#base:focus, select#sort:focus { outline: 2px solid var(--accent); outline-offset: -1px; }
  select#sort { color: var(--text-dim); cursor: pointer; }
  #settings-toggle { font: inherit; padding: 8px 10px; border: 1px solid var(--border-strong); border-radius: var(--radius);
                      background: var(--panel); color: var(--text-dim); cursor: pointer; line-height: 1; }
  #settings-toggle:hover { color: var(--text); border-color: var(--text-faint); }
  #settings-panel { display: none; align-items: center; gap: 8px; margin: -2px 0 14px; padding: 10px 12px;
                     background: var(--panel); border: 1px solid var(--border); border-radius: var(--radius); }
  #settings-panel.open { display: flex; }
  #settings-panel label { font-size: 0.8em; color: var(--text-dim); white-space: nowrap; }
  input#base { width: 220px; font-size: 0.85em; color: var(--text-dim); font-family: var(--font-mono); }
  #settings-panel .hint { font-size: 0.78em; color: var(--text-faint); }

  #stats { color: var(--text-dim); margin-bottom: 12px; font-size: 0.86em; font-variant-numeric: tabular-nums; }
  #stats b { color: var(--text); font-weight: 600; }

  .table-wrap { background: var(--panel); border: 1px solid var(--border); border-radius: var(--radius);
                overflow: auto; max-height: calc(100vh - 152px); }
  table { border-collapse: collapse; width: 100%; font-size: 0.88em; }
  th, td { text-align: left; padding: 9px 14px; border-bottom: 1px solid var(--border); white-space: nowrap; }
  td:nth-child(2) { white-space: normal; min-width: 260px; }
  th { position: sticky; top: 0; background: var(--panel); color: var(--text-dim); font-weight: 600;
       font-size: 0.82em; text-transform: uppercase; letter-spacing: .04em; border-bottom: 1px solid var(--border-strong); z-index: 1; }
  td:last-child { font-variant-numeric: tabular-nums; }
  .id-col { color: var(--text-faint); font-family: var(--font-mono); font-size: 0.92em; width: 1%;
             font-variant-numeric: tabular-nums; }
  tbody tr[data-open] { cursor: pointer; }
  tbody tr[data-open]:hover { background: var(--accent-soft); }
  tbody tr:last-child td { border-bottom: none; }
  .pill { display: inline-block; padding: 2px 8px; border-radius: 999px; font-size: 0.82em; font-weight: 600; }
  .pill-y { background: var(--ok-soft); color: var(--ok); }
  .pill-n { background: transparent; color: var(--no); }
  .tag { color: var(--text-dim); }
  .tag.muted { color: var(--text-faint); }
  .name-cell { color: var(--text); display: flex; align-items: center; gap: 6px; }
  .open-hint { opacity: 0; color: var(--accent); font-size: 0.85em; transition: opacity .1s; flex: none; }
  tr[data-open]:hover .open-hint { opacity: 1; }
  .more-row td { color: var(--text-faint); padding: 12px 14px; }
  #applied-filters { display: flex; flex-wrap: wrap; gap: 6px; margin-bottom: 10px; }
  .filter-chip { display: inline-flex; align-items: center; gap: 6px; background: var(--accent-soft); color: var(--accent-text);
                 padding: 3px 6px 3px 10px; border-radius: 999px; font-weight: 600; font-size: 0.86em; }
  .filter-chip .axis { opacity: 0.65; font-weight: 500; }
  .filter-chip button { font: inherit; border: none; background: transparent; color: inherit; cursor: pointer;
                         width: 16px; height: 16px; border-radius: 50%; display: inline-flex; align-items: center; justify-content: center; }
  .filter-chip button:hover { background: rgba(0,0,0,.12); }
  .filter-chip.clear-all { background: transparent; color: var(--text-dim); font-weight: 500; cursor: pointer; }
  .filter-chip.clear-all:hover { color: var(--text); text-decoration: underline; }

  #pager { display: flex; align-items: center; justify-content: space-between; gap: 12px; margin-top: 12px;
           font-size: 0.86em; color: var(--text-dim); }
  #pager .pages { display: flex; align-items: center; gap: 4px; }
  #pager button { font: inherit; padding: 6px 10px; border: 1px solid var(--border-strong); border-radius: 6px;
                  background: var(--panel); color: var(--text); cursor: pointer; }
  #pager button:hover:not(:disabled) { border-color: var(--accent); color: var(--accent-text); }
  #pager button:disabled { opacity: 0.4; cursor: default; }
  #pager .pg-num { font-variant-numeric: tabular-nums; padding: 0 6px; }
</style></head>
<body>
<div id="sidebar">
  <div class="brand"><span class="dot"></span>HWPX 카탈로그</div>
  <div id="axis-tabs"></div>
  <div id="axis-clear"></div>
  <div id="axis-list"></div>
</div>
<div id="main">
  <h1>HWPX 서식 카탈로그 목차</h1>
  <p class="sub">발행기관 · 문서유형 · 서식종류로 걸러보고, 행을 클릭하면 원본을 좌표 그대로 새 탭에서 연다.</p>
  <div id="toolbar">
    <input id="q" placeholder="서식명 검색…">
    <select id="sort">
      <option value="id">ID순</option>
      <option value="name">서식명 가나다순</option>
      <option value="inputDesc">입력칸 많은순</option>
    </select>
    <button id="settings-toggle" title="원본 뷰어 서버 주소 설정">⚙ 서버 설정</button>
  </div>
  <div id="settings-panel">
    <label for="base">원본 뷰어(coord_view.html) 서버 주소</label>
    <input id="base" value="http://localhost:8000">
    <span class="hint">행 클릭 시 이 주소의 /web-office/coord_view.html 로 연다</span>
  </div>
  <div id="applied-filters"></div>
  <div id="stats"></div>
  <div class="table-wrap">
    <table><thead><tr><th class="id-col">ID</th><th>서식명</th><th>발행기관</th><th>문서유형</th><th>서식종류</th>
    <th>입력가능</th><th>입력칸수</th></tr></thead><tbody id="rows"></tbody></table>
  </div>
  <div id="pager"></div>
</div>
<script>
const ITEMS = __DATA_JSON__;
const GROUPS = __GROUPS_JSON__;
const TOTAL = __TOTAL__, FILLABLE = __FILLABLE__;
const AXES = Object.keys(GROUPS);
const AXIS_FIELD = {'발행기관': 'inst', '문서유형': 'doc', '서식종류': 'kind'};
const PAGE_SIZE = 50;

// 도서관 OPAC 식 패싯 검색: 축 안에서는 다중 선택 = OR, 축 사이는 AND.
let curAxis = AXES[0];
const filters = {}; // axis -> Set(선택된 값)
for (const axis of AXES) filters[axis] = new Set();
let page = 1;
let sortMode = 'id';

function matches(item, excludeAxis) {
  for (const axis of AXES) {
    if (axis === excludeAxis) continue;
    const sel = filters[axis];
    if (sel.size && !sel.has(item[AXIS_FIELD[axis]])) return false;
  }
  const q = document.getElementById('q').value.trim().toLowerCase();
  if (q && !item.name.toLowerCase().includes(q)) return false;
  return true;
}

function facetCounts(axis) {
  const counts = new Map();
  for (const item of ITEMS) {
    if (!matches(item, axis)) continue;
    const v = item[AXIS_FIELD[axis]];
    counts.set(v, (counts.get(v) || 0) + 1);
  }
  return counts;
}

function renderSidebar() {
  const tabs = document.getElementById('axis-tabs');
  tabs.innerHTML = '';
  for (const axis of AXES) {
    const b = document.createElement('button');
    b.textContent = axis;
    b.className = axis === curAxis ? 'on' : '';
    if (filters[axis].size) {
      const n = document.createElement('span');
      n.className = 'n'; n.textContent = filters[axis].size;
      b.appendChild(n);
    }
    b.onclick = () => { curAxis = axis; renderSidebar(); };
    tabs.appendChild(b);
  }

  const clearWrap = document.getElementById('axis-clear');
  clearWrap.innerHTML = '';
  if (filters[curAxis].size) {
    const btn = document.createElement('button');
    btn.textContent = '이 축 선택 해제 (' + filters[curAxis].size + ')';
    btn.onclick = () => { filters[curAxis].clear(); page = 1; renderTable(); renderSidebar(); };
    clearWrap.appendChild(btn);
  }

  const counts = facetCounts(curAxis);
  const list = document.getElementById('axis-list');
  list.innerHTML = '';
  // 원래 카탈로그 전체 순서(빈도순)를 유지하되, 다른 필터로 0건이 된 값도
  // 목록에서 사라지지 않고 흐리게 남긴다(도서관 카탈로그 UX — 선택지 자체가
  // 사라지면 "왜 안 보이지"가 되므로).
  for (const [key] of GROUPS[curAxis]) {
    const n = counts.get(key) || 0;
    const checked = filters[curAxis].has(key);
    const el = document.createElement('label');
    el.className = 'item' + (checked ? ' checked' : '') + (n === 0 && !checked ? ' zero' : '');
    el.innerHTML = '<input type="checkbox"' + (checked ? ' checked' : '') + '><span class="lbl">' + key +
      '</span><span class="cnt">' + n.toLocaleString() + '</span>';
    el.querySelector('input').onchange = () => {
      if (filters[curAxis].has(key)) filters[curAxis].delete(key); else filters[curAxis].add(key);
      page = 1; renderTable(); renderSidebar();
    };
    list.appendChild(el);
  }
}

function cell(axis, value) {
  const muted = filters[axis].size > 0;
  return '<td class="tag' + (muted ? ' muted' : '') + '">' + value + '</td>';
}

function sortRows(rows) {
  const out = rows.slice();
  if (sortMode === 'name') out.sort((a, b) => a.name.localeCompare(b.name, 'ko'));
  else if (sortMode === 'inputDesc') out.sort((a, b) => b.inputCount - a.inputCount);
  else out.sort((a, b) => a.id - b.id);
  return out;
}

function renderAppliedFilters() {
  const wrap = document.getElementById('applied-filters');
  wrap.innerHTML = '';
  let any = false;
  for (const axis of AXES) {
    for (const key of filters[axis]) {
      any = true;
      const chip = document.createElement('span');
      chip.className = 'filter-chip';
      chip.innerHTML = '<span class="axis">' + axis + '</span><span>' + key + '</span><button title="해제">×</button>';
      chip.querySelector('button').onclick = () => {
        filters[axis].delete(key); page = 1; renderTable(); renderSidebar();
      };
      wrap.appendChild(chip);
    }
  }
  if (any) {
    const clearAll = document.createElement('span');
    clearAll.className = 'filter-chip clear-all';
    clearAll.textContent = '전체 해제';
    clearAll.onclick = () => {
      for (const axis of AXES) filters[axis].clear();
      page = 1; renderTable(); renderSidebar();
    };
    wrap.appendChild(clearAll);
  }
}

function renderPager(totalRows) {
  const pager = document.getElementById('pager');
  const totalPages = Math.max(1, Math.ceil(totalRows / PAGE_SIZE));
  if (page > totalPages) page = totalPages;
  pager.innerHTML = '';
  if (totalRows <= PAGE_SIZE) return;
  const from = (page - 1) * PAGE_SIZE + 1;
  const to = Math.min(totalRows, page * PAGE_SIZE);
  const info = document.createElement('span');
  info.textContent = from.toLocaleString() + '–' + to.toLocaleString() + ' / ' + totalRows.toLocaleString() + '건';
  pager.appendChild(info);
  const pages = document.createElement('span');
  pages.className = 'pages';
  const mk = (label, delta, disabled) => {
    const b = document.createElement('button');
    b.textContent = label; b.disabled = disabled;
    b.onclick = () => { page += delta; renderTable(); };
    return b;
  };
  pages.appendChild(mk('← 이전', -1, page <= 1));
  const num = document.createElement('span');
  num.className = 'pg-num';
  num.textContent = page + ' / ' + totalPages + '쪽';
  pages.appendChild(num);
  pages.appendChild(mk('다음 →', 1, page >= totalPages));
  pager.appendChild(pages);
}

function renderTable() {
  let rows = ITEMS.filter(r => matches(r, null));
  rows = sortRows(rows);

  renderAppliedFilters();

  document.getElementById('stats').innerHTML =
    '<b>' + rows.length.toLocaleString() + '</b>건 표시 · 전체 ' + TOTAL.toLocaleString() +
    '건 · 입력가능 ' + FILLABLE.toLocaleString() + '건';

  const start = (page - 1) * PAGE_SIZE;
  const pageRows = rows.slice(start, start + PAGE_SIZE);

  const tb = document.getElementById('rows');
  const frag = document.createDocumentFragment();
  for (const r of pageRows) {
    const tr = document.createElement('tr');
    tr.innerHTML = '<td class="id-col">' + r.id + '</td>' +
      '<td class="name-cell">' + r.name + '<span class="open-hint">↗ 원본 열기</span></td>' +
      cell('발행기관', r.inst) + cell('문서유형', r.doc) + cell('서식종류', r.kind) +
      '<td><span class="pill ' + (r.fillable ? 'pill-y' : 'pill-n') + '">' + (r.fillable ? 'Y' : 'N') + '</span></td>' +
      '<td class="tag">' + r.inputCount + '</td>';
    if (r.sourcePath) {
      tr.dataset.open = '1';
      tr.title = '클릭하면 원본을 뷰어(coord_view.html)로 연다';
      tr.onclick = () => {
        let base = document.getElementById('base').value;
        while (base.endsWith('/')) base = base.slice(0, -1);
        window.open(base + '/web-office/coord_view.html?src=' + encodeURIComponent(r.sourcePath), '_blank');
      };
    }
    frag.appendChild(tr);
  }
  tb.innerHTML = '';
  tb.appendChild(frag);
  if (!pageRows.length) {
    const tr = document.createElement('tr');
    tr.className = 'more-row';
    tr.innerHTML = '<td colspan="7">조건에 맞는 서식이 없습니다.</td>';
    tb.appendChild(tr);
  }

  renderPager(rows.length);
}

document.getElementById('q').addEventListener('input', () => { page = 1; renderTable(); renderSidebar(); });
document.getElementById('sort').addEventListener('change', (e) => { sortMode = e.target.value; renderTable(); });
document.getElementById('settings-toggle').addEventListener('click', () => {
  document.getElementById('settings-panel').classList.toggle('open');
});
renderSidebar();
renderTable();
</script>
</body></html>
""".replace("__DATA_JSON__", data_json).replace("__GROUPS_JSON__", groups_json) \
   .replace("__TOTAL__", str(total)).replace("__FILLABLE__", str(fillable_n))


if __name__ == "__main__":
    result = build()
    print(f"총 {result['total']}건 (입력가능 {result['fillable']}건)")
    print("발행기관 상위5:", result["by_inst_top5"])
    print("문서유형:", result["by_doc"])
    print("산출:", result["xlsx"])
    print("뷰어:", result["viewer"])
